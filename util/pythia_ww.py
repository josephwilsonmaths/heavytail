"""Restartable full-checkpoint WW 0.2.7 analysis and Pythia correlations."""

import gc
import json
import time
import warnings
from pathlib import Path

import numpy as np
import pandas as pd
import torch
from scipy import stats

from scripts_to_save_ww.ww_fit_metrics import fit_paper_e_tpl

DEFAULT_MODELS = tuple(f'pythia-{size}-deduped' for size in
                       ('70m', '160m', '410m', '1b', '1.4b', '2.8b', '6.9b', '12b'))
SCHEMA = 'pythia_all_matrix_weights_ww027_etpl_v1'


def _save_json(path, obj):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    temp = path.with_suffix(path.suffix + '.tmp')
    temp.write_text(json.dumps(obj, indent=2, allow_nan=False) + '\n', encoding='utf-8')
    temp.replace(path)


def summary_path(model, step=143000, output_dir='results/pythia/ww'):
    return Path(output_dir) / 'summaries' / f'{model}_step{step}_summary.json'


def load_summary(model, step=143000, output_dir='results/pythia/ww'):
    path = summary_path(model, step, output_dir)
    if not path.exists():
        raise FileNotFoundError(f'Run the full-checkpoint WW script first: missing {path}')
    saved = json.loads(path.read_text(encoding='utf-8'))
    if (saved.get('schema') != SCHEMA or saved.get('status') != 'complete'
            or saved.get('step') != step or saved.get('model') != model):
        raise ValueError(f'Invalid or mismatched cached summary: {path}')
    return saved


def _finite(value):
    return float(value) if value is not None and np.isfinite(value) else None


def measure_matrix(weight, name, watcher, *, memory_guard=True):
    """Call actual WW 0.2.7, then paper E-TPL on the unscaled full spectrum.

    WW's original randomized truncated SVD, defaults and dense scaling are
    retained. E-TPL uses the full W^T W spectrum, via the smaller Gram matrix;
    nonzero eigenvalues of W^T W and W W^T agree. Only one matrix is in RAM.
    """
    if weight.ndim != 2:
        raise ValueError('Expected a matrix weight')
    if memory_guard:
        import psutil
        rows, cols = weight.shape
        # Legacy WW also computes W.T @ W and a randomized-SVD workspace.
        estimate = 4 * (5 * rows * cols + cols * cols) + 16 * min(rows, cols) ** 2
        if psutil.virtual_memory().available < estimate:
            raise MemoryError(f'{name}: estimated workspace {estimate / 1e9:.1f} GB exceeds '
                              f'available RAM {psutil.virtual_memory().available / 1e9:.1f} GB. '
                              'Close other applications and rerun to resume.')
    started = time.perf_counter()
    W = weight.detach().cpu().to(torch.float32).numpy()
    with warnings.catch_warnings():
        warnings.simplefilter('ignore')
        res = watcher.analyze_weights(
            weights=[W], layerid=0, min_size=50, max_size=0,
            alphas=True, lognorms=True, spectralnorms=True, softranks=False,
            normalize=False, glorot_fix=False, plot=False, mp_fit=False)[0]
    keys = ('alpha', 'alpha_weighted', 'logpnorm', 'lognorm', 'logspectralnorm')
    if any(_finite(res.get(k)) is None for k in keys):
        raise RuntimeError(f'WW 0.2.7 did not return all required metrics for {name}')
    result = {'name': name, 'shape': list(weight.shape),
              **{k: _finite(res.get(k)) for k in keys}}
    # Release sklearn's temporary arrays before the separate E-TPL spectrum.
    del W, res
    gc.collect()
    with torch.no_grad():
        W64 = weight.detach().cpu().to(torch.float64)
        gram = W64.T @ W64 if W64.shape[0] >= W64.shape[1] else W64 @ W64.T
        evals = torch.linalg.eigvalsh(gram).numpy()
        del W64, gram
    tpl = fit_paper_e_tpl(evals)
    result.update(lambda_e_tpl=tpl['Lambda'], beta_e_tpl=tpl['beta'],
                  lambda_xmin=tpl['xmin'], lambda_xmax=tpl['xmax'],
                  lambda_noise_flag=tpl['noise_flag'], elapsed_sec=time.perf_counter() - started)
    return result


def summarize_matrices(records, *, model, step, revision, powerlaw_version):
    if not records:
        raise ValueError('No matrix measurements')
    result = {'model': model, 'step': int(step), 'revision': revision,
              'schema': SCHEMA, 'status': 'complete', 'ww_version': '0.2.7',
              'powerlaw_version': powerlaw_version, 'num_matrices': len(records),
              'weight_scope': 'all 2D weight tensors, including input/output embeddings; no biases/1D norms',
              'ww_method': 'original WW 0.2.7 analyze_weights defaults; equal matrix mean',
              'lambda_method': 'paper_e_tpl_v1; full unscaled spectrum; equal matrix/layer mean',
              'lambda_reference': 'https://github.com/nsfzyzz/Generalization_metrics_for_NLP',
              'lambda_num_flagged': sum(bool(r['lambda_noise_flag']) for r in records)}
    for key in ('alpha', 'alpha_weighted', 'logpnorm', 'lognorm', 'logspectralnorm', 'lambda_e_tpl'):
        result[key] = float(np.mean([r[key] for r in records]))
    result.update(log_alpha_norm=result['logpnorm'], log_norm=result['lognorm'],
                  log_spectral_norm=result['logspectralnorm'], **{'lambda': result['lambda_e_tpl']})
    return result


def run_model(model, *, step=143000, output_dir='results/pythia/ww', **kwargs):
    """Lock each model's cache so simultaneous notebook/CLI runs cannot mix it."""
    from filelock import FileLock
    model = str(model).removeprefix('EleutherAI/')
    if not model.startswith('pythia-') or '/' in model or '\\' in model:
        raise ValueError('Expected a Pythia model name')
    if summary_path(model, step, output_dir).exists():
        return load_summary(model, step, output_dir)
    locks = Path(output_dir) / 'locks'
    locks.mkdir(parents=True, exist_ok=True)
    with FileLock(locks / f'{model}_step{step}.lock', timeout=0):
        return _run_model(model, step=step, output_dir=output_dir, **kwargs)


def _run_model(model, *, step=143000, output_dir='results/pythia/ww',
              cache_dir='results/pythia/ww/checkpoint_cache', token=None,
              memory_guard=True, verbose=True):
    """Download every checkpoint shard, stream every matrix, and cache progress.

    Completed summaries are returned without importing WW or contacting HF.
    Partial results pin the original commit, so revision changes cannot mix runs.
    Downloads use the logged-in HF account or HF_TOKEN; no token is saved.
    """
    model = str(model).removeprefix('EleutherAI/')
    if not model.startswith('pythia-') or '/' in model or '\\' in model:
        raise ValueError('Expected a Pythia model name')
    path = summary_path(model, step, output_dir)
    if path.exists():
        return load_summary(model, step, output_dir)
    from huggingface_hub import HfApi, hf_hub_download
    import powerlaw
    from util.pythia_spectra import _read_weights
    from scripts_to_save_ww.ww_legacy_compat import import_legacy_weightwatcher

    ww = import_legacy_weightwatcher()
    watcher = ww.WeightWatcher(model=torch.nn.Linear(2, 2), log=False)
    out = Path(output_dir)
    partial_path = out / 'partial' / f'{model}_step{step}.json'
    partial = json.loads(partial_path.read_text()) if partial_path.exists() else None
    if partial and (partial.get('schema') != SCHEMA or
                    partial.get('powerlaw_version') != str(powerlaw.__version__)):
        raise ValueError('Partial results have different algorithm/package settings.')
    repo_id = f'EleutherAI/{model}'
    revision = partial['revision'] if partial else f'step{step}'
    info = HfApi().model_info(repo_id, revision=revision, token=token)
    files = {entry.rfilename for entry in info.siblings}
    pinned = info.sha

    def log(message):
        if verbose:
            print(f'[{model} step{step}] {message}', flush=True)

    def download(filename):
        return Path(hf_hub_download(repo_id, filename, revision=pinned,
                                   cache_dir=cache_dir, token=token))

    download('config.json')
    records = partial['records'] if partial else {}
    for index_name in ('model.safetensors.index.json', 'pytorch_model.bin.index.json'):
        if index_name in files:
            mapping = json.loads(download(index_name).read_text())['weight_map']
            shards = sorted(set(mapping.values()))
            break
    else:
        single = next((f for f in ('model.safetensors', 'pytorch_model.bin') if f in files), None)
        if single is None:
            raise FileNotFoundError(f'No supported weights for {repo_id}/{revision}')
        shards, mapping = [single], None
    expected = set(partial.get('expected', [])) if partial else set()
    log(f'{len(records)} matrices already measured; {len(shards)} checkpoint shards')
    for shard in shards:
        log(f'Downloading/reusing {shard}')
        shard_path = download(shard)
        if mapping is not None:
            names = sorted(name for name, filename in mapping.items() if filename == shard)
        elif shard_path.suffix == '.safetensors':
            from safetensors import safe_open
            with safe_open(str(shard_path), framework='pt', device='cpu') as handle:
                names = sorted(handle.keys())
        else:
            try:
                state = torch.load(shard_path, map_location='cpu', weights_only=True, mmap=True)
            except RuntimeError:
                state = torch.load(shard_path, map_location='cpu', weights_only=True)
            names = sorted(state)
            del state
        iterator = _read_weights(shard_path, [n for n in names if n.endswith('.weight')])
        try:
            for name, weight in iterator:
                if weight.ndim != 2:
                    continue
                expected.add(name)
                if name not in records:
                    log(f'Measuring {name}: {tuple(weight.shape)}')
                    records[name] = measure_matrix(weight, name, watcher, memory_guard=memory_guard)
                    _save_json(partial_path, {'model': model, 'step': step, 'schema': SCHEMA,
                                             'revision': pinned, 'powerlaw_version': str(powerlaw.__version__),
                                             'expected': sorted(expected), 'records': records})
                    log(f'Saved {len(records)} matrix measurements')
                del weight
        finally:
            iterator.close()
        gc.collect()
    if not expected or set(records) != expected:
        raise RuntimeError('Incomplete checkpoint coverage; refusing a completed summary.')
    values = [records[name] for name in sorted(records)]
    summary = summarize_matrices(values, model=model, step=step, revision=pinned,
                                 powerlaw_version=str(powerlaw.__version__))
    (out / 'details').mkdir(parents=True, exist_ok=True)
    pd.DataFrame(values).to_csv(out / 'details' / f'{model}_step{step}_details.csv', index=False)
    _save_json(path, summary)
    summaries = [json.loads(p.read_text()) for p in sorted((out / 'summaries').glob('*_summary.json'))]
    pd.DataFrame(summaries).to_csv(out / 'summary_metrics.csv', index=False)
    partial_path.unlink(missing_ok=True)
    log(f'Completed: {path}')
    return summary


def bfe_w(eigenvalues, jitter=0.0):
    """Match llm_ht: EmbedOut, lambda=1/mean(e), n=len(e), normalized BFE."""
    e = np.asarray(eigenvalues, dtype=float).reshape(-1)
    shifted = e + jitter
    if not len(e) or not np.all(np.isfinite(e)) or e.mean() <= 0 or np.any(shifted <= 0):
        raise ValueError('BFE-W requires positive finite shifted eigenvalues and a positive mean.')
    lam = 1.0 / e.mean()
    energy = lam / 2 * shifted.mean() - .5 * np.log(shifted).mean() + .5 * np.log(2 * np.pi / lam)
    return float(energy), float(lam)


def build_comparison(models=DEFAULT_MODELS, *, step=143000, output_dir='results/pythia/ww',
                     spectra_dir='results/pythia', evals_root='evals',
                     task='lambada_openai', shot='zero-shot', metric='acc', jitter=0.0):
    """Read cached WW summaries and match BFE-W/evaluation at the exact step."""
    from helpers import load_pythia_eval_metric
    rows = []
    for model in models:
        path = summary_path(model, step, output_dir)
        if not path.exists():
            raise FileNotFoundError(f'Run the full-checkpoint WW script first: missing {path}')
        ww = load_summary(model, step, output_dir)
        spectrum_path = Path(spectra_dir) / f'{model}_pile_b2m_{step}_0.pt'
        spectrum = torch.load(spectrum_path, map_location='cpu', weights_only=False)
        energy, lam = bfe_w(spectrum['EmbedOut']['eigvals'], jitter)
        evaluations = load_pythia_eval_metric(
            model_size=model.split('-')[1], task=task, metric=metric, evals_root=evals_root,
            deduped='deduped' in model, shot=shot)
        selected = evaluations.loc[evaluations['step'].eq(step), 'value']
        if len(selected) != 1:
            raise ValueError(f'Missing exact step{step} evaluation for {model}/{task}/{metric}')
        rows.append({'Model': model, 'Group': 'Deduped' if 'deduped' in model else 'Non-deduped',
                     'Step': step, 'Task': task, 'Shot': shot, 'Evaluation metric': metric,
                     'Evaluation': float(selected.iloc[0]), 'BFE-W': energy,
                     'BFE-W lambda': lam, 'WW matrices': ww['num_matrices'],
                     'Log Frobenius Norm': ww['log_norm'],
                     'Log Spectral Norm': ww['log_spectral_norm'],
                     'Weighted-alpha': ww['alpha_weighted'], 'Alpha-Norm': ww['log_alpha_norm'],
                     'Alpha (PL)': ww['alpha'], 'Lambda (E-TPL paper)': ww['lambda'],
                     'WW version': ww['ww_version'], 'Powerlaw version': ww['powerlaw_version'],
                     'Lambda flagged fits': ww['lambda_num_flagged'],
                     'Lambda reference': ww['lambda_reference'], 'Revision': ww['revision']})
    return pd.DataFrame(rows)


METRICS = ('BFE-W', 'Log Frobenius Norm', 'Log Spectral Norm', 'Weighted-alpha',
           'Alpha-Norm', 'Alpha (PL)', 'Lambda (E-TPL paper)')


def correlation_tables(data, sample_mode='common'):
    if sample_mode not in ('common', 'pairwise'):
        raise ValueError('sample_mode must be common or pairwise')
    evaluation_rows, bfe_rows = [], []
    groups = [('All models', data)]
    if data['Group'].nunique() > 1:
        groups.extend(data.groupby('Group', sort=False))
    for group_name, group in groups:
        numeric = group[['Evaluation'] + list(METRICS)].apply(pd.to_numeric, errors='coerce')
        numeric = numeric.replace([np.inf, -np.inf], np.nan)
        if sample_mode == 'common':
            numeric = numeric.dropna()
        for target, rows in [('Evaluation', evaluation_rows), ('BFE-W', bfe_rows)]:
            for key in METRICS:
                if target == key:
                    continue
                pair = numeric[[key, target]].dropna()
                coeffs = [np.nan] * 3
                if len(pair) >= 2 and pair[key].nunique() > 1 and pair[target].nunique() > 1:
                    x, y = pair[key], pair[target]
                    coeffs = [stats.pearsonr(x, y)[0], stats.spearmanr(x, y)[0], stats.kendalltau(x, y)[0]]
                rows.append({'Group': group_name, 'Metric': key, 'N': len(pair),
                             'Pearson': coeffs[0], 'Spearman': coeffs[1], 'Kendall': coeffs[2]})
    return pd.DataFrame(evaluation_rows), pd.DataFrame(bfe_rows)


def export_comparison(data, correlations, bfe_correlations, path):
    """Export numeric results and all source model rows to one Excel workbook."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(path, engine='openpyxl') as writer:
        from openpyxl.styles import Alignment
        for title, frame in [('Correlations', correlations), ('BFE vs WW', bfe_correlations),
                             ('Model metrics', data)]:
            frame.to_excel(writer, sheet_name=title, index=False, freeze_panes=(1, 0))
            ws = writer.sheets[title]
            ws.auto_filter.ref = ws.dimensions
            for idx, column in enumerate(frame.columns, 1):
                from openpyxl.utils import get_column_letter
                width = max(len(str(column)), *(len(str(v)) for v in frame[column].dropna()))
                ws.column_dimensions[get_column_letter(idx)].width = min(width + 2, 42)
                if width > 40:
                    for cells in ws.iter_rows(min_row=1, min_col=idx, max_col=idx):
                        cells[0].alignment = Alignment(wrap_text=True, vertical='center')
                        ws.row_dimensions[cells[0].row].height = max(
                            ws.row_dimensions[cells[0].row].height or 15, 30)
                if pd.api.types.is_float_dtype(frame[column]):
                    for cells in ws.iter_rows(min_row=2, min_col=idx, max_col=idx):
                        cells[0].number_format = '0.0000'
    return path
